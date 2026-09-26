from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter as L

import os
OUT = os.environ.get("OUT", os.path.join(os.path.dirname(os.path.abspath(__file__)), "SISKOP_Template_Migrasi_v1.xlsx"))
N = int(os.environ.get("N", 2000))
NR = 150          # Neraca / Produk / Penyesuaian rows
F = "Arial"
TEAL = "0E7A60"; GREY = "5F6D67"; LIGHT = "EEF2F0"; YEL = "FFF4C2"; AUTO = "E6E9E8"
thin = Side(style="thin", color="C4CDC8")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
MONEY = '#,##0.00;[Red]-#,##0.00;"-"'
DATE = "dd/mm/yyyy"

wb = Workbook()
font = lambda **k: Font(name=F, **k)

def header(ws, cols):
    """cols: list of (key, required, width, description, fmt)"""
    for i, (key, req, width, desc, fmt) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=key)
        auto = key.startswith("_")
        fill = AUTO if auto else (TEAL if req else "5F6D67")
        c.fill = PatternFill("solid", fgColor=fill)
        c.font = font(bold=True, color="1B2320" if auto else "FFFFFF", size=10)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
        c.comment = Comment(("WAJIB. " if req else ("OTOMATIS — jangan diisi. " if auto else "Opsional. ")) + desc, "SISKOP")
        ws.column_dimensions[L(i)].width = width
        for r in range(2, ws.max_row_target + 1):
            cell = ws.cell(row=r, column=i)
            cell.font = font(size=10, color="5F6D67" if auto else "000000")
            if fmt: cell.number_format = fmt
            if auto: cell.fill = PatternFill("solid", fgColor="F4F6F5")
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

def dv_list(ws, formula, rng, allow_blank=True, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=allow_blank, showErrorMessage=True,
                        errorTitle="Nilai tidak valid", error="Pilih dari daftar.")
    if prompt: dv.promptTitle, dv.prompt, dv.showInputMessage = "Petunjuk", prompt, True
    ws.add_data_validation(dv); dv.add(rng)

# ---------------------------------------------------------------- Daftar (hidden lists)
lists = {
    "A": ("jenisKoperasi", ["SYARIAH", "KONVENSIONAL"]),
    "B": ("sumberData", ["LEGACY", "EXCEL", "PDF", "CAMPURAN"]),
    "C": ("statusAnggota", ["AKTIF", "KELUAR"]),
    "D": ("yaTidak", ["Y", "N"]),
    "E": ("jenisProduk", ["SIMPANAN", "PEMBIAYAAN"]),
    "F": ("jenisSimpanan", ["POKOK", "WAJIB", "SUKARELA"]),
    "G": ("jenisPembiayaan", ["SYARIAH", "KONVENSIONAL"]),
    "H": ("tipeImbalan", ["BUNGA", "BAGI_HASIL", "MARGIN", "HARIAN"]),
    "I": ("periode", ["DAILY", "MONTHLY", "YEARLY"]),
    "J": ("kolektibilitas", ["LANCAR", "DALAM_PERHATIAN", "KURANG_LANCAR", "DIRAGUKAN", "MACET"]),
    "K": ("kategoriAkun", ["ASET", "KEWAJIBAN", "EKUITAS"]),
    "L": ("saldoNormal", ["DEBIT", "KREDIT"]),
    "M": ("kelasEkuitas", ["SIMPANAN_POKOK", "SIMPANAN_WAJIB", "MODAL_TETAP", "CADANGAN_UMUM", "CADANGAN_RISIKO",
                           "HIBAH", "MODAL_PENYERTAAN", "SHU", "EKUITAS_LAIN"]),
    "N": ("jenisUnit", ["KSP", "KONSUMEN", "PRODUSEN", "JASA", "PEMASARAN"]),
}
# ---------------------------------------------------------------- Petunjuk (filled last, created first for order)
pet = wb.active; pet.title = "Petunjuk"
info = wb.create_sheet("Info_Koperasi")
sheets = {}
for name, rows in [("Unit", 50), ("Neraca", NR + 1), ("Produk", 100), ("Anggota", N + 1), ("Simpanan", N + 1),
                   ("Pembiayaan", N + 1), ("Penyesuaian", NR + 1), ("Aset_Tetap", 300)]:
    ws = wb.create_sheet(name); ws.max_row_target = rows; sheets[name] = ws
rek = wb.create_sheet("Rekonsiliasi")
daf = wb.create_sheet("Daftar")
for col, (title, vals) in lists.items():
    daf[f"{col}1"] = title; daf[f"{col}1"].font = font(bold=True)
    for i, v in enumerate(vals, 2): daf[f"{col}{i}"] = v
daf.sheet_state = "hidden"
def lst(col): return f"=Daftar!${col}$2:${col}${len(lists[col][1]) + 1}"

TOL = "Info_Koperasi!$C$10"

# ---------------------------------------------------------------- Info_Koperasi
info.column_dimensions["A"].width = 4; info.column_dimensions["B"].width = 30
info.column_dimensions["C"].width = 42; info.column_dimensions["D"].width = 60
info["B2"] = "Info Koperasi & Batch Migrasi"; info["B2"].font = font(bold=True, size=14, color=TEAL)
info["B3"] = "Isi sel kuning. Satu file = satu batch migrasi untuk satu koperasi."; info["B3"].font = font(size=10, color=GREY)
fields = [
    ("namaKoperasi", "Nama resmi koperasi sesuai akta", "Koperasi Karyawan Contoh Sejahtera", None),
    ("subdomain", "Subdomain tenant SISKOP (tanpa .siskop...)", "kopkar-contoh", None),
    ("jenisKoperasi", "SYARIAH atau KONVENSIONAL", "SYARIAH", lst("A")),
    ("tanggalCutover", "Tanggal tutup buku yang disahkan RAT — semua saldo per tanggal ini", None, None),
    ("sumberData", "Asal data: LEGACY / EXCEL / PDF / CAMPURAN", "EXCEL", lst("B")),
    ("toleransiSelisih", "Toleransi pembulatan per akun (Rp). Default 1 — ubah hanya dengan persetujuan QA", 1, None),
    ("disiapkanOleh", "Nama staf onboarding SISKOP", "", None),
    ("tanggalDisiapkan", "Tanggal file ini disiapkan", None, None),
    ("versiBatch", "Naikkan setiap kali file dikirim ulang (v1, v2, ...)", "v1", None),
    ("dokumenSumber", "Daftar berkas asli dari klien (nama file)", "", None),
]
info["B4"], info["C4"], info["D4"] = "Kolom", "Isi", "Keterangan"
for c in ("B4", "C4", "D4"):
    info[c].font = font(bold=True, color="FFFFFF"); info[c].fill = PatternFill("solid", fgColor=TEAL)
for i, (k, d, v, dv) in enumerate(fields, 5):
    info[f"B{i}"] = k; info[f"B{i}"].font = font(bold=True)
    cell = info[f"C{i}"]; cell.value = v
    cell.fill = PatternFill("solid", fgColor=YEL); cell.border = BOX; cell.font = font(color="0000FF")
    info[f"D{i}"] = d; info[f"D{i}"].font = font(size=10, color=GREY); info[f"D{i}"].alignment = Alignment(wrap_text=True)
    if dv: dv_list(info, dv, f"C{i}")
    if k in ("tanggalCutover", "tanggalDisiapkan"): cell.number_format = DATE
assert info["B10"].value == "toleransiSelisih"
# ---------------------------------------------------------------- Unit
header(sheets["Unit"], [
    ("kodeUnit", True, 14, "Kode singkat unit, dipakai di sheet lain. Contoh: KSP01.", "@"),
    ("namaUnit", True, 34, "Nama unit usaha. Setiap koperasi minimal punya 1 unit.", None),
    ("jenisUnit", True, 16, "KSP / KONSUMEN / PRODUSEN / JASA / PEMASARAN.", None),
])
dv_list(sheets["Unit"], lst("N"), "C2:C50")
UNIT = "=Unit!$A$2:$A$50"
# ---------------------------------------------------------------- Neraca
nr = sheets["Neraca"]
header(nr, [
    ("kodeAkun", True, 12, "Kode akun COA SISKOP (mis. 1-1000). Unik.", "@"),
    ("namaAkun", True, 34, "Nama akun di SISKOP.", None),
    ("kategori", True, 13, "ASET / KEWAJIBAN / EKUITAS (neraca pembuka saja).", None),
    ("saldoNormal", True, 12, "DEBIT atau KREDIT. Akun kontra (mis. akumulasi penyusutan) = kebalikan kategorinya.", None),
    ("kelasEkuitas", False, 20, "Wajib untuk EKUITAS: dipakai untuk Modal Sendiri & Laporan Perubahan Ekuitas.", None),
    ("kodeUnit", False, 11, "Unit pemilik saldo. Kosong = tingkat koperasi (tak teralokasi).", "@"),
    ("saldo", True, 18, "Saldo per tanggal cutover, positif searah saldo normal. Angka murni, tanpa 'Rp'.", MONEY),
    ("namaDiSumber", True, 34, "Nama baris persis seperti di Neraca RAT klien — untuk ditelusuri.", None),
])
R = NR + 1
dv_list(nr, lst("K"), f"C2:C{R}"); dv_list(nr, lst("L"), f"D2:D{R}"); dv_list(nr, lst("M"), f"E2:E{R}")
dv_list(nr, UNIT, f"F2:F{R}")
AKUN = f"=Neraca!$A$2:$A${R}"
# ---------------------------------------------------------------- Produk
pr = sheets["Produk"]
header(pr, [
    ("kodeProduk", True, 13, "Kode singkat produk, dipakai di Simpanan/Pembiayaan. Contoh: SW, SSK, MRB.", "@"),
    ("namaProduk", True, 30, "Nama produk di SISKOP.", None),
    ("jenis", True, 13, "SIMPANAN atau PEMBIAYAAN.", None),
    ("jenisSimpanan", False, 14, "Untuk SIMPANAN: POKOK / WAJIB / SUKARELA.", None),
    ("jenisPembiayaan", False, 15, "Untuk PEMBIAYAAN: SYARIAH / KONVENSIONAL.", None),
    ("tipeImbalan", True, 13, "BUNGA / BAGI_HASIL / MARGIN / HARIAN.", None),
    ("tarifPersen", True, 11, "Tarif dalam persen (2.5 = 2,5%). Isi 0 jika tidak ada.", "0.00"),
    ("periode", False, 11, "Untuk SIMPANAN: DAILY / MONTHLY / YEARLY.", None),
    ("tenorMaksBulan", False, 13, "Untuk PEMBIAYAAN: tenor maksimum.", "0"),
    ("kodeAkun", True, 12, "Akun Neraca tempat saldo produk ini tercatat. Beberapa produk boleh ke satu akun.", "@"),
    ("kodeUnit", True, 11, "Unit pemilik produk.", "@"),
    ("namaDiSumber", True, 30, "Nama produk/kolom di berkas klien, mis. 'Tabungan Qurban'.", None),
])
dv_list(pr, lst("E"), "C2:C100"); dv_list(pr, lst("F"), "D2:D100"); dv_list(pr, lst("G"), "E2:E100")
dv_list(pr, lst("H"), "F2:F100"); dv_list(pr, lst("I"), "H2:H100"); dv_list(pr, AKUN, "J2:J100"); dv_list(pr, UNIT, "K2:K100")
PROD = "=Produk!$A$2:$A$100"
# ---------------------------------------------------------------- Anggota
an = sheets["Anggota"]; RN = N + 1
header(an, [
    ("noAnggotaLama", True, 14, "Nomor anggota di sistem/berkas lama. Disimpan sebagai teks. Harus unik di file ini — jika klien memakai nomor yang sama untuk 2 orang, beri akhiran (226-A, 226-B) dan catat.", "@"),
    ("namaLengkap", True, 32, "Nama sesuai KTP, gelar boleh ikut. Hapus catatan seperti '(M)', '(jukir)' ke kolom catatan.", None),
    ("nik", False, 20, "16 digit, format TEKS. Kosong = anggota ditandai data belum lengkap, KYC dilengkapi setelah go-live. JANGAN mengarang NIK.", "@"),
    ("alamat", False, 34, "Alamat KTP.", None),
    ("tempatLahir", False, 16, "", None),
    ("tanggalLahir", False, 13, "dd/mm/yyyy.", DATE),
    ("pekerjaan", False, 20, "", None),
    ("noHp", False, 15, "Format teks, mis. 081234567890.", "@"),
    ("tanggalMasuk", False, 13, "Tanggal menjadi anggota, dd/mm/yyyy.", DATE),
    ("status", True, 10, "AKTIF atau KELUAR (keluar tapi masih punya saldo).", None),
    ("pengurus", False, 10, "Y jika pengurus (batas pinjaman pihak terkait).", None),
    ("pengawas", False, 10, "Y jika pengawas.", None),
    ("kodeUnit", True, 11, "Unit tempat anggota terdaftar.", "@"),
    ("catatan", False, 30, "Catatan konversi, mis. 'nama di sheet pembiayaan: Solikin, ST MT'.", None),
    ("_cek_duplikat", False, 13, "Otomatis: GANDA jika noAnggotaLama muncul lebih dari sekali.", None),
    ("_cek_nik", False, 13, "Otomatis: OK / KOSONG / TIDAK VALID.", None),
])
dv_list(an, lst("C"), f"J2:J{RN}"); dv_list(an, lst("D"), f"K2:L{RN}"); dv_list(an, UNIT, f"M2:M{RN}")
for r in range(2, RN + 1):
    an[f"O{r}"] = f'=IF(A{r}="","",IF(COUNTIF($A$2:$A${RN},A{r})>1,"GANDA","OK"))'
    an[f"P{r}"] = f'=IF(A{r}="","",IF(C{r}="","KOSONG",IF(AND(LEN(C{r})=16,ISNUMBER(--C{r})),"OK","TIDAK VALID")))'
ANG = f"Anggota!$A$2:$A${RN}"
# ---------------------------------------------------------------- Simpanan
si = sheets["Simpanan"]
header(si, [
    ("noAnggotaLama", True, 14, "Harus ada di sheet Anggota.", "@"),
    ("namaAnggota", False, 30, "Hanya untuk pengecekan manual — sistem memakai noAnggotaLama.", None),
    ("kodeProduk", True, 12, "Kode dari sheet Produk (jenis SIMPANAN).", "@"),
    ("saldo", True, 18, "Saldo per cutover. Satu baris per anggota per produk. Tidak boleh negatif.", MONEY),
    ("catatan", False, 30, "", None),
    ("_akun", False, 11, "Otomatis: akun Neraca dari produk.", None),
    ("_cek_anggota", False, 13, "Otomatis: OK / TIDAK ADA di sheet Anggota.", None),
])
dv_list(si, PROD, f"C2:C{RN}")
for r in range(2, RN + 1):
    si[f"F{r}"] = f'=IF(C{r}="","",IFERROR(INDEX(Produk!$J$2:$J$100,MATCH(C{r},Produk!$A$2:$A$100,0)),"PRODUK?"))'
    si[f"G{r}"] = f'=IF(A{r}="","",IF(COUNTIF({ANG},A{r})=0,"TIDAK ADA","OK"))'
# ---------------------------------------------------------------- Pembiayaan
pb = sheets["Pembiayaan"]
header(pb, [
    ("noAnggotaLama", True, 14, "Harus ada di sheet Anggota.", "@"),
    ("namaAnggota", False, 28, "Hanya untuk pengecekan manual.", None),
    ("kodeProduk", True, 12, "Kode dari sheet Produk (jenis PEMBIAYAAN).", "@"),
    ("noKontrak", False, 14, "Nomor akad/kontrak bila ada. Satu baris = satu kontrak aktif.", "@"),
    ("tanggalCair", False, 13, "Tanggal pencairan asli bila diketahui.", DATE),
    ("pokokAwal", False, 16, "Plafon/pokok awal bila diketahui.", MONEY),
    ("tenorBulan", False, 10, "Tenor awal bila diketahui.", "0"),
    ("sisaPokok", True, 17, "Sisa pokok per cutover. Inilah yang dicocokkan ke Piutang di Neraca.", MONEY),
    ("sisaMarginBunga", False, 16, "Syariah: sisa margin belum diterima. Konvensional: bunga tertunggak.", MONEY),
    ("angsuranPokok", False, 15, "Angsuran pokok per bulan.", MONEY),
    ("angsuranMarginBunga", False, 17, "Angsuran margin/bunga (infaq) per bulan.", MONEY),
    ("sisaAngsuranBulan", True, 13, "Sisa jumlah bulan angsuran.", "0"),
    ("hariTunggakan", False, 12, "Hari keterlambatan per cutover. 0 jika lancar.", "0"),
    ("kolektibilitas", True, 16, "LANCAR / DALAM_PERHATIAN / KURANG_LANCAR / DIRAGUKAN / MACET.", None),
    ("kodeAkunKhusus", False, 14, "Isi hanya jika saldo tercatat di akun lain dari akun produk, mis. Piutang Macet.", "@"),
    ("catatan", False, 26, "", None),
    ("_akun", False, 11, "Otomatis: kodeAkunKhusus bila diisi, selain itu akun produk.", None),
    ("_cek_anggota", False, 13, "Otomatis: OK / TIDAK ADA di sheet Anggota.", None),
    ("_cek_angsuran", False, 14, "Otomatis: SELISIH jika sisaPokok ≠ angsuranPokok × sisaAngsuranBulan (di luar toleransi).", None),
])
dv_list(pb, PROD, f"C2:C{RN}"); dv_list(pb, lst("J"), f"N2:N{RN}"); dv_list(pb, AKUN, f"O2:O{RN}")
for r in range(2, RN + 1):
    pb[f"Q{r}"] = f'=IF(C{r}="","",IF(O{r}<>"",O{r},IFERROR(INDEX(Produk!$J$2:$J$100,MATCH(C{r},Produk!$A$2:$A$100,0)),"PRODUK?")))'
    pb[f"R{r}"] = f'=IF(A{r}="","",IF(COUNTIF({ANG},A{r})=0,"TIDAK ADA","OK"))'
    pb[f"S{r}"] = f'=IF(H{r}="","",IF(OR(J{r}="",L{r}=""),"-",IF(ABS(H{r}-J{r}*L{r})>{TOL},"SELISIH","OK")))'
# ---------------------------------------------------------------- Penyesuaian
pe = sheets["Penyesuaian"]
header(pe, [
    ("kodeAkun", True, 12, "Akun Neraca yang saldonya tidak bisa dialokasikan ke anggota.", "@"),
    ("jumlah", True, 18, "Bagian saldo Neraca yang TIDAK ada rinciannya per anggota (positif searah saldo normal; negatif jika rincian anggota lebih besar dari Neraca).", MONEY),
    ("alasan", True, 46, "Penjelasan jelas, mis. 'Rincian pembiayaan lebih besar Rp 527,5 jt dari Neraca RAT; Bendahara belum dapat merinci'.", None),
    ("buktiSumber", False, 28, "Sheet/sel atau dokumen asal selisih.", None),
    ("disetujuiOleh", True, 24, "Nama & jabatan pengurus yang menyetujui (Ketua/Bendahara).", None),
    ("tanggalPersetujuan", True, 14, "dd/mm/yyyy.", DATE),
])
dv_list(pe, AKUN, f"A2:A{NR + 1}")
# ---------------------------------------------------------------- Aset_Tetap
at = sheets["Aset_Tetap"]
header(at, [
    ("namaAset", True, 30, "Opsional untuk go-live; dipakai untuk register aset.", None),
    ("jumlahUnit", False, 10, "", "0"),
    ("tanggalPerolehan", False, 14, "", DATE),
    ("hargaPerolehan", True, 17, "", MONEY),
    ("akumulasiPenyusutan", False, 18, "Positif. Di Neraca dicatat sebagai akun kontra (saldo normal KREDIT).", MONEY),
    ("_nilaiBuku", False, 16, "Otomatis: harga perolehan − akumulasi penyusutan.", MONEY),
    ("kodeAkun", True, 12, "Akun aset tetap di Neraca.", "@"),
    ("catatan", False, 26, "", None),
])
dv_list(at, AKUN, "G2:G300")
for r in range(2, 301):
    at[f"F{r}"] = f'=IF(D{r}="","",D{r}-N(E{r}))'

# ---------------------------------------------------------------- Rekonsiliasi
rk = rek
for col, w in zip("ABCDEFGHI", [4, 12, 32, 18, 18, 18, 18, 18, 16]):
    rk.column_dimensions[col].width = w
rk["B2"] = "Rekonsiliasi — wajib hijau semua sebelum dikirim ke klien untuk persetujuan"
rk["B2"].font = font(bold=True, size=14, color=TEAL)
rk["B3"] = "Semua sel di sheet ini otomatis. Jangan diedit."; rk["B3"].font = font(size=10, color=GREY)
S = f"Simpanan!$F$2:$F${RN}"; P = f"Pembiayaan!$Q$2:$Q${RN}"
checks = [
    ("Jumlah anggota", f"=COUNTA(Anggota!$A$2:$A${RN})", None, "info"),
    ("Nomor anggota ganda", f'=COUNTIF(Anggota!$O$2:$O${RN},"GANDA")', "zero", "Harus 0"),
    ("Anggota tanpa NIK (data belum lengkap)", f'=COUNTIF(Anggota!$P$2:$P${RN},"KOSONG")', None, "Boleh > 0 — KYC dilengkapi setelah go-live"),
    ("NIK tidak valid (bukan 16 digit)", f'=COUNTIF(Anggota!$P$2:$P${RN},"TIDAK VALID")', "zero", "Harus 0 — kosongkan jika tidak yakin"),
    ("Baris simpanan tanpa anggota", f'=COUNTIF(Simpanan!$G$2:$G${RN},"TIDAK ADA")', "zero", "Harus 0"),
    ("Baris simpanan dengan produk tak dikenal", f'=COUNTIF({S},"PRODUK?")', "zero", "Harus 0"),
    ("Saldo simpanan negatif", f'=COUNTIF(Simpanan!$D$2:$D${RN},"<0")', "zero", "Harus 0"),
    ("Pembiayaan tanpa anggota", f'=COUNTIF(Pembiayaan!$R$2:$R${RN},"TIDAK ADA")', "zero", "Harus 0"),
    ("Pembiayaan dengan produk tak dikenal", f'=COUNTIF({P},"PRODUK?")', "zero", "Harus 0"),
    ("Pembiayaan: sisa pokok ≠ angsuran × sisa bulan", f'=COUNTIF(Pembiayaan!$S$2:$S${RN},"SELISIH")', None, "Periksa; boleh > 0 bila ada pelunasan sebagian"),
    ("Total saldo normal DEBIT (Neraca)", f'=SUMIFS(Neraca!$G$2:$G${R},Neraca!$D$2:$D${R},"DEBIT")', None, "info"),
    ("Total saldo normal KREDIT (Neraca)", f'=SUMIFS(Neraca!$G$2:$G${R},Neraca!$D$2:$D${R},"KREDIT")', None, "info"),
    ("Saldo Awal Migrasi (DEBIT − KREDIT)", "=D16-D17", "tol", "Harus 0 (± toleransi)"),
    ("Akun Neraca yang belum cocok", '=COUNTIF($I$26:$I$175,"SELISIH")', "zero", "Harus 0"),
    ("Jumlah akun Neraca terisi", f'=COUNTA(Neraca!$A$2:$A${R})', "pos", "Harus > 0"),
    ("Tanggal cutover terisi", '=IF(Info_Koperasi!$C$8="",0,1)', "pos", "Harus terisi di Info_Koperasi"),
    ("Penyesuaian belum disetujui pengurus", f'=SUMPRODUCT((Penyesuaian!$A$2:$A${NR + 1}<>"")*(Penyesuaian!$E$2:$E${NR + 1}=""))', "zero", "Harus 0 — isi disetujuiOleh setelah konfirmasi klien"),
]
rk["B5"], rk["C5"], rk["D5"], rk["E5"], rk["F5"] = "", "Pemeriksaan", "Hasil", "Status", "Syarat"
for c in ("C5", "D5", "E5", "F5"):
    rk[c].font = font(bold=True, color="FFFFFF"); rk[c].fill = PatternFill("solid", fgColor=TEAL)
for i, (label, formula, kind, note) in enumerate(checks, 6):
    rk[f"C{i}"] = label; rk[f"C{i}"].font = font(size=10)
    rk[f"D{i}"] = formula; rk[f"D{i}"].font = font(size=10)
    rk[f"D{i}"].number_format = MONEY if "Total" in label or "Saldo Awal" in label else "#,##0"
    if kind == "zero": rk[f"E{i}"] = f'=IF(D{i}=0,"OK","PERBAIKI")'
    elif kind == "pos": rk[f"E{i}"] = f'=IF(D{i}>0,"OK","PERBAIKI")'
    elif kind == "tol": rk[f"E{i}"] = f'=IF(ABS(D{i})<={TOL},"OK","PERBAIKI")'
    else: rk[f"E{i}"] = '="-"'
    rk[f"F{i}"] = note; rk[f"F{i}"].font = font(size=10, color=GREY)
    rk[f"E{i}"].font = font(size=10, bold=True)
assert rk["C16"].value.startswith("Total saldo normal DEBIT") and rk["C18"].value.startswith("Saldo Awal")
rk["C24"] = "STATUS BATCH"; rk["C24"].font = font(bold=True, size=12)
rk["D24"] = '=IF(COUNTIF(E6:E22,"PERBAIKI")=0,"SIAP DISETUJUI KLIEN","BELUM SIAP")'
rk["D24"].font = font(bold=True, size=12)
rk.merge_cells("D24:F24")

rk["B24"] = "Per akun Neraca: saldo Neraca = rincian anggota + penyesuaian"; rk["B24"].font = font(bold=True, size=12, color=TEAL)
heads = ["kodeAkun", "namaAkun", "Saldo Neraca", "Rincian simpanan", "Rincian pembiayaan", "Penyesuaian", "Selisih", "Status"]
for j, h in enumerate(heads):
    c = rk.cell(row=25, column=2 + j, value=h)
    c.font = font(bold=True, color="FFFFFF", size=10); c.fill = PatternFill("solid", fgColor=TEAL)
    c.alignment = Alignment(horizontal="center", wrap_text=True)
for k in range(NR):
    r = 26 + k; src = 2 + k
    rk[f"B{r}"] = f'=IF(Neraca!A{src}="","",Neraca!A{src})'
    rk[f"C{r}"] = f'=IF(B{r}="","",Neraca!B{src})'
    rk[f"D{r}"] = f'=IF(B{r}="","",N(Neraca!G{src}))'
    rk[f"E{r}"] = f'=IF(B{r}="","",SUMIFS(Simpanan!$D$2:$D${RN},{S},B{r}))'
    rk[f"F{r}"] = f'=IF(B{r}="","",SUMIFS(Pembiayaan!$H$2:$H${RN},{P},B{r}))'
    rk[f"G{r}"] = f'=IF(B{r}="","",SUMIFS(Penyesuaian!$B$2:$B${NR + 1},Penyesuaian!$A$2:$A${NR + 1},B{r}))'
    backed = (f'COUNTIF(Produk!$J$2:$J$100,B{r})+COUNTIF(Pembiayaan!$O$2:$O${RN},B{r})'
              f'+COUNTIF(Penyesuaian!$A$2:$A${NR + 1},B{r})')
    rk[f"H{r}"] = f'=IF(B{r}="","",IF({backed}=0,"",D{r}-E{r}-F{r}-G{r}))'
    rk[f"I{r}"] = f'=IF(B{r}="","",IF({backed}=0,"tanpa rincian",IF(ABS(H{r})<={TOL},"COCOK","SELISIH")))'
    for col in "DEFGH": rk[f"{col}{r}"].number_format = MONEY
    for col in "BCDEFGHI": rk[f"{col}{r}"].font = font(size=10)
rk.freeze_panes = "A6"
green = PatternFill("solid", fgColor="D8F0E6"); red = PatternFill("solid", fgColor="F9D7D3")
rk.conditional_formatting.add("E6:E22", CellIsRule(operator="equal", formula=['"OK"'], fill=green))
rk.conditional_formatting.add("E6:E22", CellIsRule(operator="equal", formula=['"PERBAIKI"'], fill=red))
rk.conditional_formatting.add("I26:I175", CellIsRule(operator="equal", formula=['"COCOK"'], fill=green))
rk.conditional_formatting.add("I26:I175", CellIsRule(operator="equal", formula=['"SELISIH"'], fill=red))
rk.conditional_formatting.add("D24", CellIsRule(operator="equal", formula=['"SIAP DISETUJUI KLIEN"'], fill=green))
rk.conditional_formatting.add("D24", CellIsRule(operator="equal", formula=['"BELUM SIAP"'], fill=red))
for ws, rng in [(an, f"O2:P{RN}"), (si, f"F2:G{RN}"), (pb, f"Q2:S{RN}")]:
    for bad in ("GANDA", "TIDAK VALID", "TIDAK ADA", "PRODUK?", "SELISIH"):
        ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{bad}"'], fill=red))

# ---------------------------------------------------------------- Petunjuk
pet.column_dimensions["A"].width = 3
for col, w in zip("BCDEFGH", [16, 22, 9, 12, 26, 58, 4]): pet.column_dimensions[col].width = w
pet["B2"] = "SISKOP — Template Migrasi Saldo Awal v1"; pet["B2"].font = font(bold=True, size=16, color=TEAL)
pet["B3"] = "Untuk tim onboarding SISKOP. Berkas klien (Excel/PDF/ekspor sistem lama) dikonversi ke template ini, lalu diunggah sebagai draf batch migrasi di tenant klien."
pet["B3"].font = font(size=10, color=GREY)
r = 5
def para(title, lines):
    global r
    pet[f"B{r}"] = title; pet[f"B{r}"].font = font(bold=True, size=12, color=TEAL); r += 1
    for ln in lines:
        pet[f"B{r}"] = ln; pet[f"B{r}"].font = font(size=10); pet[f"B{r}"].alignment = Alignment(wrap_text=False); r += 1
    r += 1
para("Alur kerja", [
    "1. Klien mengirim dokumen: Neraca yang disahkan RAT, daftar anggota, saldo simpanan per anggota, daftar pembiayaan aktif.",
    "2. Tim SISKOP mengisi template ini — urutan: Info_Koperasi → Unit → Neraca → Produk → Anggota → Simpanan → Pembiayaan → Penyesuaian.",
    "3. Buka sheet Rekonsiliasi. Semua status harus OK dan STATUS BATCH = 'SIAP DISETUJUI KLIEN'. Selisih ditanyakan ke Bendahara klien.",
    "4. Unggah ke tenant klien sebagai DRAF batch migrasi (belum aktif). Sistem menjalankan pemeriksaan yang sama.",
    "5. Ketua & Bendahara klien meninjau laporan di aplikasi dan menyetujui → saldo awal diposting per tanggal cutover.",
])
para("Aturan pengisian", [
    "• Jangan mengubah nama sheet atau judul kolom baris 1 — sistem membacanya persis. Kolom berawalan '_' dan sheet Rekonsiliasi terisi otomatis.",
    "• Angka murni tanpa 'Rp', titik ribuan, atau spasi. Desimal memakai titik/koma sesuai Excel Anda; maksimal 2 desimal.",
    "• Tanggal dd/mm/yyyy. Semua saldo per tanggalCutover di Info_Koperasi.",
    "• NIK, noAnggotaLama, noHp, noKontrak berformat TEKS (agar nol di depan tidak hilang). Jangan mengarang NIK — kosongkan bila tidak ada.",
    "• Satu baris per anggota per produk (Simpanan) dan satu baris per kontrak aktif (Pembiayaan).",
    "• Baris non-anggota (mis. 'Kopkar', 'bunga admin bank') TIDAK dimasukkan ke Simpanan — saldo itu ikut Neraca atau dijelaskan di Penyesuaian.",
    "• Jangan memaksa angka cocok. Selisih yang disetujui klien dicatat di Penyesuaian dengan alasan dan nama penyetuju.",
    "• Jangan membetulkan kesalahan pembukuan klien secara diam-diam (mis. penyusutan dijumlah ke aktiva). Tanyakan dan catat.",
    "• Ekspor CSV (bila diperlukan): satu CSV per sheet, UTF-8, baris 1 = judul kolom, kolom berawalan '_' boleh ikut — diabaikan sistem.",
])
para("Keterangan warna", [])
legend = [(TEAL, "FFFFFF", "Judul kolom WAJIB"), ("5F6D67", "FFFFFF", "Judul kolom opsional"),
          (AUTO, "1B2320", "Kolom otomatis (rumus) — jangan diisi"), (YEL, "0000FF", "Sel isian di Info_Koperasi")]
r -= 1
for fill, fc, text in legend:
    c = pet[f"B{r}"]; c.value = "contoh"; c.fill = PatternFill("solid", fgColor=fill); c.font = font(size=10, color=fc, bold=True)
    pet[f"C{r}"] = text; pet[f"C{r}"].font = font(size=10); r += 1
r += 1
pet[f"B{r}"] = "Kamus kolom (dengan contoh pengisian)"; pet[f"B{r}"].font = font(bold=True, size=12, color=TEAL); r += 1
kh = ["Sheet", "Kolom", "Wajib", "Tipe", "Contoh", "Keterangan"]
for j, h in enumerate(kh):
    c = pet.cell(row=r, column=2 + j, value=h); c.font = font(bold=True, color="FFFFFF", size=10)
    c.fill = PatternFill("solid", fgColor=TEAL)
r += 1
examples = {
    "Unit": ["KSP01", "Simpan Pinjam Syariah", "KSP"],
    "Neraca": ["2-1100", "Simpanan Sukarela Anggota", "KEWAJIBAN", "KREDIT", "", "", "3841964000", "Simpanan Sukarela"],
    "Produk": ["SSK", "Tabungan Sukarela", "SIMPANAN", "SUKARELA", "", "BAGI_HASIL", "0", "YEARLY", "", "2-1100", "KSP01", "Tabungan Sukarela"],
    "Anggota": ["118", "Ety Sukmayati, BA", "3578xxxxxxxxxxxx", "Jl. Sutorejo No. 59, Surabaya", "Surabaya", "12/05/1968", "Karyawan", "081234567890", "01/03/1998", "AKTIF", "N", "N", "KSP01", "", "(otomatis)", "(otomatis)"],
    "Simpanan": ["118", "Ety Sukmayati, BA", "SW", "5371000", "", "(otomatis)", "(otomatis)"],
    "Pembiayaan": ["170", "Drs. Asror", "MRB", "", "", "", "", "47222600", "15583900", "1388900", "458350", "34", "0", "LANCAR", "", "", "(otomatis)", "(otomatis)", "(otomatis)"],
    "Penyesuaian": ["1-1100", "-527501000", "Rincian pembiayaan per anggota lebih besar dari Neraca RAT", "Pembiayaan!L174", "Sujiati S.W., Bendahara", "15/01/2026"],
    "Aset_Tetap": ["Printer", "1", "01/06/2021", "1000000", "600000", "(otomatis)", "1-2100", ""],
}
for name in ["Unit", "Neraca", "Produk", "Anggota", "Simpanan", "Pembiayaan", "Penyesuaian", "Aset_Tetap"]:
    ws = sheets[name]
    for j in range(1, ws.max_column + 1):
        key = ws.cell(row=1, column=j).value
        cm = ws.cell(row=1, column=j).comment.text
        req = "Ya" if cm.startswith("WAJIB") else ("Auto" if cm.startswith("OTOMATIS") else "")
        fmt = ws.cell(row=2, column=j).number_format
        tipe = {"@": "Teks", DATE: "Tanggal", MONEY: "Rupiah", "0": "Angka", "0.00": "Persen"}.get(fmt, "Teks")
        desc = cm.split(". ", 1)[1] if ". " in cm else cm
        desc = desc.replace("OTOMATIS — jangan diisi. ", "")
        ex = examples[name][j - 1] if j - 1 < len(examples[name]) else ""
        vals = [name, key, req, tipe, ex, desc]
        for k, v in enumerate(vals):
            c = pet.cell(row=r, column=2 + k, value=v); c.font = font(size=9, italic=(k == 4), color=("5F6D67" if k == 4 else "000000"))
            c.alignment = Alignment(wrap_text=(k == 5), vertical="top")
            if k == 4: c.number_format = "@"
        r += 1
pet[f"B{r+1}"] = "Contoh diambil dari berkas RAT 2025 Kopkar UM Surabaya kecuali NIK/alamat (fiktif, disamarkan). Dokumen desain: docs/research/2026-09-26-client-data-migration-research.md"
pet[f"B{r+1}"].font = font(size=9, italic=True, color=GREY)

# tab colors
for name, color in [("Petunjuk", "5F6D67"), ("Info_Koperasi", "D08A2E"), ("Rekonsiliasi", TEAL)]:
    wb[name].sheet_properties.tabColor = color
for ws in wb.worksheets:
    ws.sheet_view.zoomScale = 100
wb.save(OUT)
print("saved")
